"""Supplier tables are reviewed as exact workshop costs, never catalog facts."""

import csv
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO, StringIO
import unicodedata
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from defusedxml.ElementTree import fromstring
from defusedxml.common import DefusedXmlException
from django.core import signing
from django.db import transaction
from openpyxl import load_workbook

from authentication.errors import contract_error
from pricing.repository import json_text, one, rows
from pricing.xlsx_import import MAX_BYTES, MAX_EXPANDED_BYTES, MAX_ROWS, import_rows

FIELDS = ("sku", "description", "unit", "unit_cost")
ALIASES = {
    "sku": {"sku", "codigo", "cod", "articulo", "codigo articulo"},
    "description": {"descripcion", "description", "detalle", "nombre", "producto"},
    "unit": {"unidad", "unit", "un", "unidad de compra"},
    "unit_cost": {"precio", "cost", "costo", "precio unitario", "precio neto", "valor", "unit cost"},
}
UNITS = {"BAR": "BAR", "BARRA": "BAR", "BARRAS": "BAR", "M": "M", "METRO": "M",
    "METROS": "M", "M2": "M2", "M²": "M2", "KIT": "KIT", "UNIT": "UNIT", "EA": "EA",
    "UNIDAD": "UNIT", "UNIDADES": "UNIT", "UN": "UNIT"}
SALT = "dekopen.supplier-prices.v1"


def invalid(detail):
    return contract_error(422, "supplier_prices_invalid", detail)


def _xlsx(content):
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise invalid("La planilla descomprimida supera 30 MB.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheet = workbook.active
        if (sheet is None or sheet.max_row is None or sheet.max_column is None
                or sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 100):
            raise invalid("La hoja activa admite hasta 5.000 filas y 100 columnas.")
        with ZipFile(BytesIO(content)) as archive:
            root = fromstring(archive.read(sheet._worksheet_path))
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        numbers = {cell.attrib["r"]: cell.findtext("s:v", "", ns)
            for cell in root.findall(".//s:c", ns) if cell.attrib.get("t", "n") == "n"}
        table = []
        for cells in sheet.iter_rows():
            if any(cell.data_type in ("f", "e") for cell in cells):
                raise invalid(f"Fila {cells[0].row}: reemplaza las fórmulas o errores por valores revisados.")
            # Never stringify openpyxl's binary numeric value for a price or code.
            table.append([("", False) if cell.value is None else (numbers.get(cell.coordinate, ""), True) if cell.data_type == "n"
                else (str(cell.value) if cell.value is not None else "", False) for cell in cells])
        return sheet.title, table
    except (BadZipFile, KeyError, ValueError, OSError, DefusedXmlException, ParseError) as error:
        raise invalid("No se pudo leer el Excel. Usa un archivo XLSX válido.") from error
    finally:
        if workbook is not None:
            workbook.close()


def read_table(content, file_name):
    if not content or len(content) > MAX_BYTES:
        raise invalid("El archivo de precios debe pesar como máximo 5 MB.")
    suffix = file_name.rsplit(".", 1)[-1].lower()
    if suffix == "xlsx":
        title, table = _xlsx(content)
    elif suffix == "csv":
        try:
            text = content.decode("utf-8-sig")
            dialect = csv.Sniffer().sniff(text[:10_000], delimiters=",;\t")
            table = [[(value, False) for value in row] for row in csv.reader(StringIO(text), dialect, strict=True)]
        except (UnicodeDecodeError, csv.Error) as error:
            raise invalid("Usa un CSV UTF-8 con columnas separadas por coma, punto y coma o tabulación.") from error
        title = file_name
    else:
        raise invalid("Elige una lista de precios XLSX o CSV.")
    if not table or len(table) > MAX_ROWS + 1 or any(len(row) > 100 for row in table):
        raise invalid("El archivo admite hasta 5.000 filas y 100 columnas.")
    headers = [value.strip() for value, _ in table[0]]
    if not any(headers):
        raise invalid("La primera fila debe contener los nombres de las columnas.")
    return title, headers, [(number, row) for number, row in enumerate(table[1:], 2)
        if any(value.strip() for value, _ in row)]


def _header(value):
    return " ".join("".join(char for char in unicodedata.normalize("NFKD", value.casefold())
        if not unicodedata.combining(char)).split())


def inspect(content, file_name):
    title, headers, data = read_table(content, file_name)
    suggested = {}
    for field, aliases in ALIASES.items():
        matches = [value for value in headers if _header(value) in aliases]
        suggested[field] = matches[0] if len(matches) == 1 else ""
    return {"sheet": title, "columns": headers, "suggested_mapping": suggested, "total": len(data),
        "preview": [[value for value, _ in row] for _, row in data[:10]]}


def parse(content, file_name, mapping, separator):
    _, headers, data = read_table(content, file_name)
    if (not isinstance(mapping, dict) or set(mapping) != set(FIELDS)
            or not all(isinstance(value, str) for value in mapping.values()) or len(set(mapping.values())) != 4
            or any(not value or headers.count(value) != 1 for value in mapping.values())):
        raise invalid("Elige una columna distinta y sin encabezados repetidos para cada campo.")
    if separator not in (".", ","):
        raise invalid("Elige el separador decimal de los precios de texto.")
    indices = {field: headers.index(name) for field, name in mapping.items()}
    result, seen = [], set()
    for number, row in data:
        selected = {field: row[index] if index < len(row) else ("", False) for field, index in indices.items()}
        sku, description, unit = (selected[field][0].strip() for field in FIELDS[:3])
        unit = UNITS.get(unit.upper())
        if not sku or len(sku) > 100 or sku in seen:
            raise invalid(f"Fila {number}: falta el código, supera 100 caracteres o está repetido.")
        if unit is None:
            raise invalid(f"Fila {number}: elige Barra, Metro, m², Kit o Unidad. No convertimos kilos en metros.")
        if len(description) > 1000:
            raise invalid(f"Fila {number}: la descripción supera 1.000 caracteres.")
        raw, numeric = selected["unit_cost"]
        text = raw.strip()
        if not numeric and (" " in text or ("." if separator == "," else ",") in text):
            raise invalid(f"Fila {number}: el precio mezcla separadores. Quita los separadores de miles.")
        try:
            price = Decimal(text if numeric else text.replace(",", "."))
        except InvalidOperation as error:
            raise invalid(f"Fila {number}: completa un precio decimal válido.") from error
        if not price.is_finite() or price < 0 or price >= Decimal("10000000000") or price.as_tuple().exponent < -4:
            raise invalid(f"Fila {number}: el precio admite hasta cuatro decimales y debe ser mayor o igual que cero.")
        if any(value.lstrip().startswith("=") for value, _ in row):
            raise invalid(f"Fila {number}: reemplaza las fórmulas por valores revisados.")
        result.append({"sku": sku, "description": description, "unit": unit, "unit_cost": price})
        seen.add(sku)
    if not result:
        raise invalid("El archivo no tiene filas de precios.")
    return result


def review(org_id, content, file_name, data):
    parsed = parse(content, file_name, data["mapping"], data["decimal_separator"])
    with transaction.atomic():
        cost_list = one("SELECT * FROM public.cost_lists WHERE org_id=%s AND id=%s FOR UPDATE",
            [org_id, data["cost_list_id"]], "cost_list_not_found")
        existing = rows("SELECT sku FROM public.cost_list_items WHERE org_id=%s AND cost_list_id=%s",
            [org_id, data["cost_list_id"]])
        duplicates = {item["sku"] for item in existing} & {item["sku"] for item in parsed}
        if duplicates:
            raise invalid("La lista elegida ya tiene estos códigos: " + ", ".join(sorted(duplicates)[:10]) +
                ". Elige una lista nueva o edita los precios existentes en la ficha.")
        payload = {"org": str(org_id), "list": str(data["cost_list_id"]),
            "file": sha256(content).hexdigest(), "name": file_name, "mapping": data["mapping"],
            "separator": data["decimal_separator"], "reason": data["reason"],
            "list_hash": sha256(json_text(cost_list).encode()).hexdigest()}
        if data["apply"]:
            try:
                accepted = signing.loads(data.get("review_token", ""), salt=SALT, max_age=1800)
            except signing.BadSignature as error:
                raise contract_error(409, "supplier_prices_review_required", "Vuelve a previsualizar los precios antes de guardarlos.") from error
            if accepted != payload:
                raise contract_error(409, "supplier_prices_review_changed", "Cambió el archivo, el mapeo o la lista. Revisa una vista previa nueva.")
            reason = f"{data['reason']} · {file_name[:120]} · SHA256 {payload['file']}"
            import_rows(org_id, data["cost_list_id"], parsed, reason, True)
        return {"items": [{**item, "unit_cost": str(item["unit_cost"]), "origin": "WORKSHOP"} for item in parsed[:10]],
            "total": len(parsed), "applied": data["apply"],
            "review_token": None if data["apply"] else signing.dumps(payload, salt=SALT, compress=True)}

"""Bounded XLSX parsing: no formulas, guessed locales or silent overwrites."""

from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import fromstring
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook

from dekopen_engine.commercial import PricingError
from pricing.repository import admin_write, one

MAX_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 30 * 1024 * 1024
MAX_ROWS = 5000


def parse_xlsx(content: bytes, mapping: dict[str,str], separator: str):
    if not content or len(content) > MAX_BYTES or separator not in ('.',','):
        raise PricingError('invalid_xlsx_file')
    if (not isinstance(mapping,dict) or set(mapping) != {'sku','description','unit','unit_cost'}
            or not all(isinstance(value,str) and value.strip() for value in mapping.values())
            or len(set(mapping.values())) != 4):
        raise PricingError('invalid_column_mapping')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise PricingError('xlsx_expanded_limit')
        workbook = load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
    except (BadZipFile,KeyError,ValueError,OSError,DefusedXmlException,ParseError) as error:
        raise PricingError('invalid_xlsx_file') from error
    try:
        sheet = workbook.active
        if (sheet is None or sheet.max_row is None or sheet.max_column is None
                or sheet.max_row > MAX_ROWS+1 or sheet.max_column > 100):
            raise PricingError('xlsx_sheet_limit')
        # Read monetary numeric lexemes directly: openpyxl's binary numeric
        # representation must never round a four-decimal supplier authority.
        with ZipFile(BytesIO(content)) as archive:
            root = fromstring(archive.read(sheet._worksheet_path))
        ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        numeric_text = {cell.attrib['r']:cell.find('s:v',ns).text
                        for cell in root.findall('.//s:c',ns)
                        if cell.attrib.get('t','n')=='n' and cell.find('s:v',ns) is not None}
        iterator = sheet.iter_rows()
        headers = [str(cell.value).strip() if cell.value is not None else '' for cell in next(iterator)]
        if any(headers.count(name) != 1 for name in mapping.values()):
            raise PricingError('xlsx_header_missing_or_duplicate')
        indices = {field:headers.index(name) for field,name in mapping.items()}
        result, seen = [], set()
        for row_number, cells in enumerate(iterator,2):
            if row_number > MAX_ROWS+1 or len(cells)>100:
                raise PricingError('xlsx_sheet_limit')
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.data_type in ('f','e') for cell in cells):
                raise PricingError(f'xlsx_formula_or_error_row_{row_number}')
            values = {field:cells[index].value for field,index in indices.items()}
            sku = str(values['sku']).strip() if values['sku'] is not None else ''
            unit = str(values['unit']).strip().upper()
            if not sku or len(sku)>100 or sku in seen or unit not in ('BAR','M','M2','KIT','UNIT'):
                raise PricingError(f'xlsx_invalid_or_duplicate_row_{row_number}')
            raw = numeric_text.get(cells[indices['unit_cost']].coordinate, values['unit_cost'])
            if isinstance(raw,bool) or raw is None:
                raise PricingError(f'xlsx_invalid_price_row_{row_number}')
            # Numeric cells originate in the workbook, not monetary arithmetic.
            # A textual cell requires the explicitly selected decimal separator.
            text = str(raw).strip()
            if isinstance(raw,str) and cells[indices['unit_cost']].coordinate not in numeric_text:
                if (' ' in text or (separator==',' and '.' in text)
                        or (separator=='.' and ',' in text)):
                    raise PricingError(f'xlsx_ambiguous_price_row_{row_number}')
                text = text.replace(',','.')
            try:
                price = Decimal(text)
            except InvalidOperation as error:
                raise PricingError(f'xlsx_invalid_price_row_{row_number}') from error
            if (not price.is_finite() or price < 0 or price >= Decimal('10000000000')
                    or price.as_tuple().exponent < -4):
                raise PricingError(f'xlsx_invalid_price_row_{row_number}')
            seen.add(sku)
            result.append({'sku':sku,'description':str(values['description'] or ''),
                           'unit':unit,'unit_cost':price})
        if not result:
            raise PricingError('xlsx_no_rows')
        return result
    finally:
        workbook.close()


def import_rows(org_id, cost_list_id, parsed, reason, apply):
    one('SELECT id FROM public.cost_lists WHERE id=%s AND org_id=%s FOR UPDATE',
        [cost_list_id,org_id],'cost_list_not_found')
    if not apply:
        return parsed
    # Duplicate existing SKU is an explicit constraint failure; never implicit upsert.
    for item in parsed:
        admin_write('cost-items',org_id,{'cost_list_id':cost_list_id,'sku':item['sku'],
                    'unit':item['unit'],'unit_cost':item['unit_cost'],'description':item['description'],
                    'item_type':'IMPORTED'},reason)
    return parsed

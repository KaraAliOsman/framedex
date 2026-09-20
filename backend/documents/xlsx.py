"""Deterministic exact-text XLSX representations for frozen supplier orders."""

from __future__ import annotations

from copy import copy
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.writer.excel import ExcelWriter

from documents.repository import DocumentaryError

_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _object(value: object, code: str) -> dict[str, object]:
    if not isinstance(value, dict) or not all(isinstance(key, str) for key in value):
        raise DocumentaryError(code)
    return value


def _array(value: object, code: str) -> list[object]:
    if not isinstance(value, list):
        raise DocumentaryError(code)
    return value


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, (str, int)):
        return str(value)
    raise DocumentaryError("xlsx_value_not_exact_text")


def _normalize_archive(content: bytes) -> bytes:
    source = BytesIO(content)
    target = BytesIO()
    with ZipFile(source, "r") as incoming, ZipFile(target, "w", ZIP_DEFLATED) as outgoing:
        for name in sorted(incoming.namelist()):
            original = incoming.getinfo(name)
            info = ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = original.compress_type
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            outgoing.writestr(info, incoming.read(name))
    return target.getvalue()


def _verify_matrix(content: bytes, sheet_name: str, expected: list[list[str | int]]) -> None:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name]
        for index, row in enumerate(sheet.iter_rows(max_row=len(expected))):
            expected_row = expected[index]
            actual = row[: len(expected_row)]
            if len(actual) != len(expected_row):
                raise DocumentaryError("xlsx_frozen_value_mismatch")
            for cell, wanted in zip(actual, expected_row, strict=True):
                if cell.data_type == "f":
                    raise DocumentaryError("xlsx_formula_cell_forbidden")
                actual_value = "" if cell.value is None else cell.value
                actual_type = "s" if cell.value is None else cell.data_type
                wanted_type = (
                    "n"
                    if isinstance(wanted, int) and not isinstance(wanted, bool)
                    else "s"
                )
                if actual_value != wanted or actual_type != wanted_type:
                    raise DocumentaryError("xlsx_frozen_value_mismatch")
    finally:
        workbook.close()


def render_order_xlsx(document_type: str, snapshot: dict[str, object]) -> tuple[bytes, str]:
    order = _object(snapshot.get("order"), "invalid_order_snapshot")
    revision = _object(snapshot.get("revision"), "invalid_order_snapshot")
    lines = [_object(item, "invalid_order_line")
             for item in _array(snapshot.get("lines"), "invalid_order_snapshot")]
    if document_type == "DOC-02" and order.get("order_type") != "SUPPLIER_GLASS_PO":
        raise DocumentaryError("document_scope_mismatch")
    if document_type == "DOC-04" and order.get("order_type") != "SUPPLIER_PROFILE_PO":
        raise DocumentaryError("document_scope_mismatch")
    if document_type not in ("DOC-02", "DOC-04"):
        raise DocumentaryError("xlsx_document_type_invalid")

    if document_type == "DOC-02":
        sheet_name = "Pedido de vidrios"
        headers = [
            "Requisito", "SKU técnico", "SKU compra", "Composición", "Ancho mm",
            "Alto mm", "Cantidad", "Unidad", "Pulido", "Ubicación", "Trazabilidad",
            "Área m²",
        ]
        data = []
        total_area = Decimal("0")
        for line in lines:
            specification = _object(line.get("specification"), "invalid_order_line")
            polishing = _object(specification.get("polishing"), "invalid_order_line")
            quantity = int(line["quantity"])
            area = (
                Decimal(_text(specification.get("oriented_width_mm")))
                * Decimal(_text(specification.get("oriented_height_mm")))
                * quantity / Decimal("1000000")
            )
            total_area += area
            data.append([
                _text(line.get("requirement_key")),
                ", ".join(_text(item) for item in _array(
                    line.get("technical_skus"), "invalid_order_line"
                )),
                _text(line.get("purchasing_sku")),
                _text(specification.get("composition")),
                _text(specification.get("oriented_width_mm")),
                _text(specification.get("oriented_height_mm")),
                quantity,
                _text(line.get("unit")),
                "/".join(
                    edge.upper() for edge in ("top", "right", "bottom", "left")
                    if polishing.get(edge) is True
                ) or "SIN PULIDO",
                _text(specification.get("location_tag")),
                ", ".join(_text(item) for item in _array(
                    line.get("source_trace"), "invalid_order_line"
                )),
                format(area.normalize(), "f"),
            ])
        data.append([
            "TOTAL", "—", "—", "—", "—", "—", "—", "—", "—", "—", "—",
            format(total_area.normalize(), "f"),
        ])
    else:
        sheet_name = "Pedido de perfiles"
        headers = [
            "Requisito", "Categoría", "SKU taller", "SKU compra", "Stock físico",
            "Largo barra mm", "Cantidad", "Unidad", "Perfil de corte", "Trazabilidad",
        ]
        data = []
        for line in lines:
            specification = _object(line.get("specification"), "invalid_order_line")
            data.append([
                _text(line.get("requirement_key")),
                _text(line.get("category")),
                ", ".join(_text(item) for item in _array(
                    line.get("technical_skus"), "invalid_order_line"
                )),
                _text(line.get("purchasing_sku")),
                _text(line.get("physical_stock_identity")),
                _text(specification.get("stock_length_mm")),
                int(line["quantity"]),
                _text(line.get("unit")),
                _text(specification.get("cutting_profile_id")),
                ", ".join(_text(item) for item in _array(
                    line.get("source_trace"), "invalid_order_line"
                )),
            ])

    metadata = [
        ["Documento", document_type],
        ["Orden", _text(order.get("order_code"))],
        ["Proveedor", _text(order.get("supplier_name"))],
        ["Revisión", _text(revision.get("revision_code"))],
        ["BOM hash", _text(revision.get("bom_hash"))],
        ["Snapshot hash", _text(revision.get("snapshot_sha256"))],
    ]
    expected: list[list[str | int]] = [*metadata, headers, *data]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row_index, row in enumerate(expected, start=1):
        for column_index, value in enumerate(row, start=1):
            if isinstance(value, bool) or not isinstance(value, (int, str)):
                raise DocumentaryError("xlsx_value_not_exact_text")
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if isinstance(value, str):
                # Frozen authority text stays a literal string even when it
                # begins with =, +, - or @; openpyxl would otherwise store a
                # leading '=' as an executable formula.
                cell.data_type = "s"
    sheet.freeze_panes = "A8"
    last_data_row = len(expected) - (1 if document_type == "DOC-02" else 0)
    sheet.auto_filter.ref = f"A7:{chr(64 + len(headers))}{last_data_row}"
    for cell in sheet[7]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    for column in sheet.columns:
        width = min(60, max(12, max(len(_text(cell.value)) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc).replace(tzinfo=None)
    workbook.properties.creator = "Dekopen"
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    raw = BytesIO()
    with ZipFile(raw, "w", ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(workbook, archive).write_data()
    workbook.close()
    content = _normalize_archive(raw.getvalue())
    _verify_matrix(content, sheet_name, expected)
    return content, _XLSX_MEDIA

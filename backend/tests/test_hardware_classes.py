"""D04 catalog interchange, price preview and sealed downstream projections."""

from contextlib import nullcontext
from copy import deepcopy
from decimal import Decimal as D
from io import BytesIO, StringIO
import csv
import json
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from catalogs.demo_hardware import hardware_manifest
from catalogs.serializers import KitWriteSerializer
from dekopen_engine.geometry import calculate_geometry
from engine.tests.hardware_cases import hardware_case_inputs
from ingest.catalog_template import candidate, export_csv, parse_structured
from inventory.production_stock import unit_stock_needs
from portal.service import _sealed_positions
from production.service import _routing, version_hardware_picking, _require_hardware_machining_authority
from projects import hardware


def imported_kit():
    kit = hardware_manifest()[0]["params"]["available_hardware_kits"][0]
    return {**kit,"system_code":"DEMO_60","is_active":True,"source":"Fuente sintética de prueba"}


def test_class_csv_and_filled_template_roundtrip_without_decimal_loss():
    entry = candidate("Herrajes",imported_kit(),key="kit",row=5,method="MANUAL")
    assert not entry["errors"]
    content = export_csv("Herrajes",[entry])
    parsed = parse_structured("CSV",content)
    assert parsed[0]["values"] == entry["values"]
    workbook = load_workbook("backend/ingest/templates/catalogo-v1.xlsx")
    row = list(csv.reader(StringIO(content.decode("utf-8-sig")),delimiter=";"))[2]
    for column,value in enumerate(row,1):
        workbook["Herrajes"].cell(5,column,value)
    stream = BytesIO()
    workbook.save(stream)
    assert parse_structured("XLSX",stream.getvalue())[0]["values"] == entry["values"]
    fields = KitWriteSerializer().fields
    data = {name:entry["values"].get(name) for name in fields if name!="system_id"}
    serializer = KitWriteSerializer(data={**data,"system_id":str(uuid4())})
    assert serializer.is_valid(),serializer.errors
    assert serializer.validated_data["class_authority"]["components"][0]["quantity"]["base"] == 2


@pytest.mark.parametrize("mutation", ["float","duplicate","no_source","two_sources","bad_count","rc_without_source"])
def test_authority_rejects_uncertified_or_ambiguous_catalog_data(mutation):
    data = imported_kit()
    authority = data["class_authority"]
    if mutation=="float":
        authority["components"][0]["weight_kg"]=0.15
    elif mutation=="duplicate":
        authority["components"].append(deepcopy(authority["components"][0]))
    elif mutation=="no_source":
        authority["source"]=""
    elif mutation=="two_sources":
        data["contents"]=[{"sku":"OLD","name":"Duplicado","qty":"1","unit":"unidad","category":"OTHER"}]
    elif mutation=="bad_count":
        authority["components"][0]["quantity"]={"base":2,"axis":"HEIGHT"}
    elif mutation=="rc_without_source":
        authority["options"][0].update(security_rating="RC2",source="")
    serializer=KitWriteSerializer(data={key:value for key,value in data.items() if key in KitWriteSerializer().fields} | {"system_id":str(uuid4())})
    assert not serializer.is_valid()


def test_sealed_components_reserve_each_length_and_never_the_kit(monkeypatch):
    _,params,node,_ = hardware_case_inputs()[0]
    result = calculate_geometry(node,params)
    monkeypatch.setattr("inventory.production_stock._mapping_rows",lambda *args,**kwargs:{})
    needs,unmapped = unit_stock_needs(org_id=uuid4(),system_id=str(uuid4()),quantity=12,
        hardware_items=[item.model_dump(mode="json") for item in result.hardware_items],fittings=[],panels=[],sheet_purchases=[])
    assert not unmapped
    assert all(row["kind"]=="FITTING" for row in needs)
    rod = next(row for row in needs if "Cremona" in row["name"])
    assert rod["variant_key"].startswith("LENGTH:")
    assert rod["needed"]==D(12)


def test_undeclared_coordinates_still_claim_a_station():
    _,params,node,_ = hardware_case_inputs()[0]
    engine = calculate_geometry(node,params).model_dump(mode="json")
    stations = _routing(engine,profile={"stations":[{"code":"CUT"},{"code":"MACHINING"},{"code":"QC","when":"required"}]})
    assert stations==["CUT","MACHINING","QC"]
    assert "MACHINING" in _routing(engine,profile=None)


def test_preview_next_class_has_motor_delta_and_no_apply_side_effect(monkeypatch):
    _,params,node,rates = hardware_case_inputs()[0]
    node=node.model_copy(update={"width_mm":D(1500),"height_mm":D(2300),"glass_spec":"6-12-6",
        "hardware_set_sku":next(kit.sku for kit in params.available_hardware_kits if kit.opening_type=="TILT_TURN" and kit.class_authority.priority==0)})
    monkeypatch.setattr(hardware.SystemParamsRepository,"load_visible",lambda *args:params)
    monkeypatch.setattr(hardware,"commercial_backend",nullcontext)
    monkeypatch.setattr(hardware,"rows",lambda *args:[{"default_margin_pct":D("0.35"),"waste_factor_pct":D("0.05")}])
    monkeypatch.setattr(hardware.PricingRepository,"cost",lambda self,sku,unit:rates[(sku,unit)])
    output=hardware.preview_hardware(uuid4(),{"system_id":uuid4(),"parametric_tree":node.model_dump(mode="json"),
        "width_mm":node.width_mm,"height_mm":node.height_mm,"bay_id":node.id,"color":"WHITE"})
    leaf=output["leaves"][0]
    assert leaf["message"] and leaf["recommendation_sku"]
    selected=next(row for row in leaf["candidates"] if row["selected"])
    recommended=next(row for row in leaf["candidates"] if row["sku"]==leaf["recommendation_sku"])
    assert not selected["compatible"] and recommended["compatible"]
    assert D(recommended["delta_net"])==D(recommended["price_net"])-D(selected["price_net"])
    assert node.hardware_set_sku==selected["sku"]


def test_revision_picking_equals_twelve_frozen_positions(monkeypatch):
    positions=[]
    expected=D(0)
    for index in range(12):
        _,params,node,_=hardware_case_inputs()[index%4]
        result=calculate_geometry(node,params)
        positions.append({"id":str(uuid4()),"position_index":index+1,"quantity":2,"engine_result":result.model_dump(mode="json")})
        expected += sum(c.qty*kit.qty*2 for kit in result.hardware_items for c in kit.contents)
    bom=[{**position,"position_id":position["id"]} for position in positions]
    monkeypatch.setattr("production.service.one",lambda *args:{"snapshot_json":json.dumps({"positions":positions,"bom":bom})})
    monkeypatch.setattr("production.service.documentary_backend",nullcontext)
    output=version_hardware_picking(org_id=uuid4(),version_id=uuid4())
    assert sum(D(row["quantity"]) for row in output["rows"])==expected
    assert all(isinstance(row["quantity"],str) for row in output["rows"])


def test_portal_discloses_only_sold_hardware():
    source={"handle_name":"Con llave","handle_color":"Negra","options":["Limitador"],"synthetic":False,"components":[{"sku":"PRIVATE"}],"cost":"991"}
    positions=_sealed_positions({"snapshot_json":json.dumps({"positions":[{"commercial_hardware":[source]}]})})
    assert positions[0]["commercial_hardware"]==[{key:source[key] for key in ("handle_name","handle_color","options","synthetic")}]


def test_unemitted_hardware_prevents_machining_completion(monkeypatch):
    from documents.repository import DocumentaryError
    from engine.tests.hardware_cases import manufacturing_case
    _,params,node,_ = hardware_case_inputs()[0]
    result = calculate_geometry(node,params)
    facts = manufacturing_case(node,params)
    snapshot = {"bom":[{"position_id":facts.position_id,"engine_result":result.model_dump(mode="json")}],
        "manufacturing":[facts.model_dump(mode="json")]}
    monkeypatch.setattr("production.service.one",lambda *args:{"snapshot_json":json.dumps(snapshot)})
    with pytest.raises(DocumentaryError,match="step_ops_incomplete") as caught:
        _require_hardware_machining_authority(org_id=uuid4(),order={"project_version_id":uuid4(),
            "payload_json":{"position_id":facts.position_id}})
    assert caught.value.extra["hardware_machining"]

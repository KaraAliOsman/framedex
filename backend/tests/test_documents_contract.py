from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from dekopen_engine.documentary_canonical import file_sha256
from documents.artifacts import _require_document_role
from documents.renderers import _doc01, _doc03, _doc06, _doc07, render_pdf_document
from documents.repository import DocumentaryError
from documents.serializers import HandleIntentSerializer
from documents.storage import SIGNED_URL_TTL_SECONDS, SupabaseDocumentStorage
from documents.xlsx import render_order_xlsx


def revision_snapshot() -> dict[str, object]:
    return {
        "schema_version": 1,
        "canonical_version": "DOCUMENTARY_CANONICAL_V1",
        "revision": "REV-A",
        "sealed_at": "2026-09-14T12:00:00Z",
        "bom_hash": "a" * 64,
        "production_allowed": True,
        "documentary_complete": True,
        "project": {
            "code": "P-001",
            "client_name": "Cliente <Seguro>",
            "delivery_address": "Obra Norte",
            "currency": "CLP",
            "total_price_net": "100000",
            "total_price_tax": "19000",
            "total_price_gross": "119000",
            "payment_terms": "50% anticipo, 50% contra entrega",
            "quotation_valid_until": "2026-10-14",
            "notes_commercial": "Incluye instalación",
        },
        "positions": [{
            "id": "11111111-1111-1111-1111-111111111111",
            "position_index": 1,
            "location_tag": "FACHADA-NORTE",
            "width_mm": "1000.00",
            "height_mm": "1200.00",
            "quantity": 2,
            "color_interior": "WHITE",
            "color_exterior": "WHITE",
            "parametric_tree": {
                "id": "B1", "type": "BAY", "opening_type": "FIXED",
                "glass_spec": "4-12-4 Float Incoloro",
                "children": [],
            },
            "workshop_annotations": [{
                "bay_id": "B1",
                "leaf_id": None,
                "bottom_drain_holes_mm": ["100.00", "500.00", "900.00"],
                "closing_points_perimeter_mm": ["150.00"],
                "continuous_width_mm": "1000.00",
                "finish_class": "WHITE",
                "has_coupler": False,
            }],
        }],
        "inspector": [{"config": {"R10": {"tolerance_mm": "1.50"}}}],
        "pricing": {
            "input_snapshot": {"cost_lines": [[1, "60000.00"]]},
            "applied_total_cost_net": "60000.00",
        },
        "realized_waste": {"status": "NOT_RECORDED", "value": None},
        "manufacturing": [],
        "purchase_requirements": {"stock_groups": []},
    }


def order_snapshot(order_type: str) -> dict[str, object]:
    glass = order_type == "SUPPLIER_GLASS_PO"
    return {
        "order": {
            "order_code": "PO-ABC123",
            "order_type": order_type,
            "supplier_name": "Proveedor Exacto",
            "project_code": "P-001",
            "confirmed_at": "2026-09-14T12:00:00Z",
        },
        "revision": {
            "revision_code": "REV-A",
            "bom_hash": "a" * 64,
            "snapshot_sha256": "b" * 64,
        },
        "lines": [{
            "requirement_key": "c" * 64,
            "category": "GLASS" if glass else "PROFILE",
            "technical_skus": ["TECH-01"],
            "purchasing_sku": "BUY-01",
            "physical_stock_identity": None if glass else "STOCK-01",
            "quantity": 2,
            "unit": "EA" if glass else "BAR",
            "source_trace": ["d" * 64, "e" * 64],
            "specification": ({
                "composition": "4-12-4 Float Incoloro",
                "oriented_width_mm": "876.00",
                "oriented_height_mm": "1076.00",
                "polishing": {"top": True, "right": False, "bottom": False, "left": True},
                "location_tag": "FACHADA-NORTE",
            } if glass else {
                "stock_length_mm": "5800.00",
                "cutting_profile_id": "CUT-PROLINE",
            }),
        }],
    }


def test_additive_bom_normalizer_keeps_old_snapshots_comparable() -> None:
    from documents.service import _without_additive_bom_fields

    current = {
        "profile_cuts": [
            {
                "sku": "MARCO-60",
                "role": "FRAME",
                "length_mm": "1000",
                "bay_id": "B1",
                "leaf_id": None,
                "sagitta_mm": None,
            }
        ],
        "reinforcements": [],
        "glasses": [
            {
                "bay_id": "B1",
                "leaf_id": None,
                "width_mm": "900",
                "height_mm": "1900",
                "thickness_net_mm": "24",
                "weight_kg": "10.26",
                "shape": None,
                "area_m2": None,
                "exposed_edges": None,
                "glass_spec": "4-16-4",
                "article_sku": "DVH-4-16-4",
            }
        ],
        "fittings": [],
    }
    legacy_stored = {
        "profile_cuts": [
            {
                "sku": "MARCO-60",
                "role": "FRAME",
                "length_mm": "1000",
                "bay_id": "B1",
                "leaf_id": None,
            }
        ],
        "reinforcements": [],
        "glasses": [
            {"bay_id": "B1", "leaf_id": None, "width_mm": "900", "height_mm": "1900"}
        ],
    }
    # Fields the model gained after the snapshot was sealed drop out on both
    # sides — the unchanged position stays comparable. Both sides normalize
    # against the same older snapshot.
    assert _without_additive_bom_fields(current, legacy_stored) == _without_additive_bom_fields(
        legacy_stored, legacy_stored
    )

    # A value carried by the snapshot stays compared: a changed bend must flag.
    bent_stored = {
        **legacy_stored,
        "profile_cuts": [
            {**legacy_stored["profile_cuts"][0], "sagitta_mm": "500"}
        ],
    }
    bent_current = {
        **current,
        "profile_cuts": [{**current["profile_cuts"][0], "sagitta_mm": "300"}],
    }
    assert _without_additive_bom_fields(bent_current, bent_stored) != _without_additive_bom_fields(
        bent_stored, bent_stored
    )

    # A null in the snapshot is equivalent to the field never having existed —
    # a non-null recomputed value must not flag drift.
    null_stored = {
        **legacy_stored,
        "glasses": [
            {
                **legacy_stored["glasses"][0],
                "glass_spec": None,
                "article_sku": None,
            }
        ],
    }
    assert _without_additive_bom_fields(current, null_stored) == _without_additive_bom_fields(
        null_stored, null_stored
    )

    # Same-role cuts of different lengths are distinct pieces: a drift on one
    # must not be masked by the other's last-wins identity row.
    double_stored = {
        **legacy_stored,
        "profile_cuts": [
            {**legacy_stored["profile_cuts"][0], "sagitta_mm": "500"},
            {
                **legacy_stored["profile_cuts"][0],
                "length_mm": "1400",
                "sagitta_mm": "300",
            },
        ],
    }
    double_current = {
        **current,
        "profile_cuts": [
            {**current["profile_cuts"][0], "sagitta_mm": "500"},
            {
                **current["profile_cuts"][0],
                "length_mm": "1400",
                "sagitta_mm": "620",
            },
        ],
    }
    assert _without_additive_bom_fields(double_current, double_stored) != _without_additive_bom_fields(
        double_stored, double_stored
    )


def test_era_projections_reproduce_historical_preimages() -> None:
    from documents.service import _drop_bom_keys

    current = {
        "profile_cuts": [{"sku": "MARCO-60", "sagitta_mm": None}],
        "reinforcements": [{"parent_profile_sku": "MARCO-60", "sagitta_mm": None}],
        "glasses": [
            {
                "bay_id": "B1",
                "glass_spec": "4-16-4",
                "article_sku": "DVH",
                "shape": None,
                "exposed_edges": None,
            }
        ],
        "fittings": [],
    }
    era94 = _drop_bom_keys(
        current, frozenset({"fittings"}), {"glasses": frozenset({"exposed_edges"})}
    )
    assert "fittings" not in era94
    assert "exposed_edges" not in era94["glasses"][0]
    assert era94["glasses"][0]["shape"] is None  # era-92 fields survive
    era92 = _drop_bom_keys(
        era94,
        frozenset(),
        {
            "glasses": frozenset({"shape"}),
            "profile_cuts": frozenset({"sagitta_mm"}),
            "reinforcements": frozenset({"sagitta_mm"}),
        },
    )
    assert "shape" not in era92["glasses"][0]
    assert "sagitta_mm" not in era92["profile_cuts"][0]
    assert "sagitta_mm" not in era92["reinforcements"][0]
    era86 = _drop_bom_keys(
        era92, frozenset(), {"glasses": frozenset({"glass_spec", "article_sku"})}
    )
    assert "glass_spec" not in era86["glasses"][0]
    assert "article_sku" not in era86["glasses"][0]


def test_client_document_escapes_input_and_never_contains_raw_cost() -> None:
    html = _doc01(revision_snapshot())
    assert "Cliente &lt;Seguro&gt;" in html
    assert "60000.00" not in html
    assert "$\u00a0119.000" in html


def test_client_quote_includes_deterministic_opening_drawings() -> None:
    html = _doc01(revision_snapshot())
    assert "<svg" in html and 'viewBox="0 0 1000 1200"' in html
    sliding = revision_snapshot()
    sliding["positions"][0]["parametric_tree"] = {  # type: ignore[index]
        "id": "B1", "type": "BAY", "opening_type": "SLIDING_2L",
        "glass_spec": "4-12-4 Float Incoloro", "children": [],
    }
    sliding_html = _doc01(sliding)
    assert sliding_html.count('marker-end="url(#arrow-1)"') == 2
    assert _doc01(sliding) == sliding_html


def test_client_quote_draws_stacked_assembly_as_a_column() -> None:
    """A door + transom STACKED assembly draws the transom ABOVE its column —
    a 1000×2200 door carrying a 1000×400 transom spans 1000×2600 with a
    horizontal seam at the contact, not a 2000-wide side-by-side row."""
    snapshot = revision_snapshot()
    snapshot["positions"][0]["parametric_tree"] = {  # type: ignore[index]
        "version": "product-v2",
        "assembly": {
            "modules": [
                {
                    "id": "door", "width_mm": "1000.00", "height_mm": "2200.00",
                    "tree": {"id": "B1", "type": "BAY", "opening_type": "FIXED",
                             "glass_spec": "4-12-4 Float Incoloro", "children": []},
                },
                {
                    "id": "transom", "width_mm": "1000.00", "height_mm": "400.00",
                    "tree": {"id": "B2", "type": "BAY", "opening_type": "FIXED",
                             "glass_spec": "4-12-4 Float Incoloro", "children": []},
                },
            ],
            "couplings": [{
                "id": "c1", "kind": "STACKED", "modules": ["door", "transom"],
                "edges": ["top", "bottom"],
            }],
        },
    }
    html = _doc01(snapshot)
    assert 'viewBox="0 0 1000 2600"' in html
    # The transom sill sits at 2200 mm elevation → svg y = 2600 − 2200 = 400.
    assert 'x1="0" y1="400" x2="1000" y2="400"' in html


def test_workshop_order_prints_annotations_drawing_and_assembly_matrix() -> None:
    snapshot = revision_snapshot()
    snapshot["manufacturing"] = [{  # type: ignore[index]
        "position_id": "11111111-1111-1111-1111-111111111111",
        "position_index": 1,
        "repetition_index": 1,
        "nominal_width_mm": "1000.00",
        "nominal_height_mm": "1200.00",
        "members": [
            {
                "member_id": "a" * 64, "bay_id": "B1",
                "identity": {"role": "FRAME", "physical_member_slot": "OUTER_LEFT"},
                "workshop_sku": "P-101", "cut_length_mm": "1200.00",
                "angle_left": "90.00", "angle_right": "90.00",
                "start": {"x_mm": "0.00", "y_mm": "0.00"},
                "end": {"x_mm": "0.00", "y_mm": "1200.00"},
            },
            {
                "member_id": "d" * 64, "bay_id": "B1",
                "identity": {"role": "FRAME", "physical_member_slot": "OUTER_RIGHT"},
                "workshop_sku": "P-101", "cut_length_mm": "1200.00",
                "angle_left": "90.00", "angle_right": "90.00",
                "start": {"x_mm": "1000.00", "y_mm": "0.00"},
                "end": {"x_mm": "1000.00", "y_mm": "1200.00"},
            },
        ],
        "reinforcements": [{"reinforcement_id": "c" * 64, "parent_member_id": "d" * 64}],
        "infills": [{
            "infill_id": "e" * 64, "bay_id": "B1", "leaf_id": "L1",
            "rect": {"width_mm": "900.00", "height_mm": "1100.00"},
        }],
        "leaves": [{"leaf_fact_id": "b" * 64, "bay_id": "B1", "leaf_id": "L1"}],
        "handles": [{
            "handle_id": "f" * 64,
            "bay_id": "B1",
            "leaf_id": "L1",
            "handle_domain_slot": "PRIMARY",
            "host_member_id": "9" * 64,
            "point": {"x_mm": "60.00", "y_mm": "1050.00"},
            "requested_height_mm": "1050.00",
            "vertical_reference": "OUTER_BOTTOM",
        }],
        "relationships": [
            {"relationship": "BELONGS_TO_LEAF", "source_id": "a" * 64, "target_id": "b" * 64},
            {"relationship": "REINFORCES", "source_id": "c" * 64, "target_id": "d" * 64},
            {"relationship": "RETAINS_INFILL", "source_id": "a" * 64, "target_id": "e" * 64},
        ],
    }]
    html = _doc03(snapshot)
    assert "100.00, 500.00, 900.00" in html
    assert "150.00" in html
    assert "Matriz de ensamble" in html
    assert "BELONGS_TO_LEAF" in html and "REINFORCES" in html
    # heterogeneous endpoints resolve to the printed piece codes, not raw ids
    assert "M-01" in html and "R-01" in html and "I-01" in html and "H-01" in html
    matrix = html.split("Matriz de ensamble", 1)[1]
    assert "a" * 64 not in matrix and "b" * 64 not in matrix
    assert "c" * 64 not in matrix and "e" * 64 not in matrix
    assert "1050.00" in html
    assert "<svg" in html


def test_qc_is_blank_and_cost_report_uses_frozen_not_recorded_authority() -> None:
    qc = _doc06(revision_snapshot())
    assert "Diferencia ≤ 1.50 mm" in qc
    assert "________________" in qc
    cost = _doc07(revision_snapshot())
    assert "$\u00a060.000" in cost
    assert "NO REGISTRADA" in cost
    assert "valor: —" in cost
    invalid = {**revision_snapshot(), "realized_waste": {"status": "RECORDED", "value": "0"}}
    with pytest.raises(DocumentaryError, match="realized_waste_authority_invalid"):
        _doc07(invalid)


def test_pdf_producer_emits_concrete_file_with_distinct_hash(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("WEASYPRINT_DLL_DIRECTORIES", "C:\\msys64\\ucrt64\\bin")
    content, media_type = render_pdf_document(
        "DOC-01", revision_snapshot(), pdf_identifier="b" * 64
    )
    assert content.startswith(b"%PDF-")
    assert media_type == "application/pdf"
    assert file_sha256(content) not in {"a" * 64, "b" * 64}


def test_xlsx_is_deterministic_exact_text_and_no_formula_authority(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class TickingDateTime:
        calls = 0

        @classmethod
        def now(cls, tz=None):
            cls.calls += 1
            return datetime(2026, 9, 20, tzinfo=tz) + timedelta(seconds=cls.calls)

    monkeypatch.setattr(
        "openpyxl.writer.excel.datetime",
        SimpleNamespace(datetime=TickingDateTime, timezone=timezone),
    )
    first, media_type = render_order_xlsx("DOC-02", order_snapshot("SUPPLIER_GLASS_PO"))
    second, _ = render_order_xlsx("DOC-02", order_snapshot("SUPPLIER_GLASS_PO"))
    assert first == second
    assert media_type.endswith("sheet")
    workbook = load_workbook(BytesIO(first), read_only=True, data_only=False)
    try:
        sheet = workbook["Pedido de vidrios"]
        assert sheet["E8"].value == "876.00"
        assert sheet["F8"].value == "1076.00"
        assert sheet["G8"].value == 2
        assert sheet["I8"].value == "TOP/LEFT"
        assert sheet["L8"].value == "1.885152"
        assert sheet["A9"].value == "TOTAL"
        assert sheet["L9"].value == "1.885152"
        assert not any(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for row in sheet.iter_rows() for cell in row
        )
    finally:
        workbook.close()
    profile, _ = render_order_xlsx("DOC-04", order_snapshot("SUPPLIER_PROFILE_PO"))
    profile_workbook = load_workbook(BytesIO(profile), read_only=True)
    try:
        assert profile_workbook["Pedido de perfiles"]["F8"].value == "5800.00"
    finally:
        profile_workbook.close()


def test_xlsx_formula_like_text_stays_literal_never_a_formula() -> None:
    snapshot = order_snapshot("SUPPLIER_GLASS_PO")
    snapshot["order"]["supplier_name"] = "=1+1"  # type: ignore[index]
    snapshot["order"]["order_code"] = "+SUM(A1:A2)"  # type: ignore[index]
    line = snapshot["lines"][0]  # type: ignore[index]
    line["purchasing_sku"] = "-1+2"
    line["technical_skus"] = ["@SUM(A1:A2)", "=cmd|' /C calc'!A0"]
    line["source_trace"] = ['=HYPERLINK("http://x")', "+POW(2,3)"]
    line["specification"]["composition"] = "@FILTER(A:A)"
    line["specification"]["location_tag"] = "-FACHADA"
    content, _ = render_order_xlsx("DOC-02", snapshot)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook["Pedido de vidrios"]
        assert sheet["B2"].value == "+SUM(A1:A2)"
        assert sheet["B3"].value == "=1+1"
        assert sheet["B8"].value == "@SUM(A1:A2), =cmd|' /C calc'!A0"
        assert sheet["C8"].value == "-1+2"
        assert sheet["D8"].value == "@FILTER(A:A)"
        assert sheet["J8"].value == "-FACHADA"
        assert sheet["K8"].value == '=HYPERLINK("http://x"), +POW(2,3)'
        assert all(
            cell.data_type != "f" for row in sheet.iter_rows() for cell in row
        )
    finally:
        workbook.close()
    profile_snapshot = order_snapshot("SUPPLIER_PROFILE_PO")
    profile_line = profile_snapshot["lines"][0]  # type: ignore[index]
    profile_line["requirement_key"] = "=REQ"
    profile_line["category"] = "@PROFILE"
    profile_line["purchasing_sku"] = "-BUY"
    profile_line["physical_stock_identity"] = "=STOCK+1"
    profile_line["specification"]["cutting_profile_id"] = "+CUT"
    profile, _ = render_order_xlsx("DOC-04", profile_snapshot)
    profile_workbook = load_workbook(BytesIO(profile), read_only=True, data_only=False)
    try:
        sheet = profile_workbook["Pedido de perfiles"]
        assert sheet["A8"].value == "=REQ"
        assert sheet["B8"].value == "@PROFILE"
        assert sheet["D8"].value == "-BUY"
        assert sheet["E8"].value == "=STOCK+1"
        assert sheet["I8"].value == "+CUT"
        assert all(
            cell.data_type != "f" for row in sheet.iter_rows() for cell in row
        )
    finally:
        profile_workbook.close()


def test_xlsx_empty_text_cells_roundtrip_as_empty() -> None:
    snapshot = order_snapshot("SUPPLIER_PROFILE_PO")
    snapshot["lines"][0]["physical_stock_identity"] = None  # type: ignore[index]
    snapshot["lines"][0]["technical_skus"] = []  # type: ignore[index]
    content, _ = render_order_xlsx("DOC-04", snapshot)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = workbook["Pedido de perfiles"]
        assert sheet["C8"].value in (None, "")
        assert sheet["E8"].value in (None, "")
        assert all(
            cell.data_type != "f" for row in sheet.iter_rows() for cell in row
        )
    finally:
        workbook.close()


def test_handle_transport_rejects_derived_manufacturing_fields() -> None:
    serializer = HandleIntentSerializer(data={
        "bay_id": "B1",
        "handle_domain_slot": "PRIMARY",
        "requested_height_mm": "1000.00",
        "vertical_reference": "OUTER_BOTTOM",
        "manufacturing_x_mm": "50.00",
    })
    assert not serializer.is_valid()


def test_document_role_matrix_keeps_doc07_owner_only() -> None:
    _require_document_role("DOC-07", "OWNER")
    for role in ("ESTIMATOR", "WORKSHOP_MANAGER", "INSTALLER"):
        with pytest.raises(DocumentaryError, match="document_access_denied"):
            _require_document_role("DOC-07", role)
    _require_document_role("DOC-02", "WORKSHOP_MANAGER")
    _require_document_role("DOC-01", "ESTIMATOR")


class FakeResponse:
    status_code = 200
    signed_url_path = "/object/sign/documents/key?token=transient"

    def json(self):
        return {"signedURL": self.signed_url_path}


class FakeClient:
    request_json: dict[str, object] | None = None

    def __init__(self, **kwargs):
        self.kwargs = kwargs

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return None

    def post(self, url, **kwargs):
        FakeClient.request_json = kwargs.get("json")
        return FakeResponse()


def test_signed_url_uses_fixed_ttl_and_is_not_artifact_identity(
    settings, monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings.SUPABASE_URL = "http://127.0.0.1:25321"
    settings.SUPABASE_SERVICE_ROLE_KEY = "local-test-service-key"
    settings.SUPABASE_STORAGE_BUCKET_DOCS = "documents"
    monkeypatch.setattr("documents.storage.httpx.Client", FakeClient)
    for path in (
        "/object/sign/documents/key?token=transient",
        "/storage/v1/object/sign/documents/key?token=transient",
    ):
        FakeResponse.signed_url_path = path
        signed = SupabaseDocumentStorage().signed_url("org_x/projects/p/rev/doc.pdf")
        assert FakeClient.request_json == {"expiresIn": SIGNED_URL_TTL_SECONDS}
        assert signed == (
            "http://127.0.0.1:25321/storage/v1/object/sign/documents/key"
            "?token=transient"
        )

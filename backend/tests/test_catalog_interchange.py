"""Exact interchange and evidence boundaries independent of the database."""
from copy import deepcopy
from pathlib import Path
import json
from uuid import uuid4

import pytest
from authentication.errors import ContractAPIException
from backend.tests.catalog_interchange_fixture import catalog_rows
from ingest import catalog_review
from ingest.catalog_ai import parse_response
from ingest.catalog_template import SCHEMAS, candidate, export_csv, normalize, parse_structured

@pytest.mark.parametrize("sheet", list(SCHEMAS))
def test_nine_sheet_csv_roundtrip_with_evidence(sheet):
    original = [row for row in catalog_rows() if row["sheet"] == sheet]
    assert original and not any(row["errors"] for row in original)
    parsed = parse_structured("CSV", export_csv(sheet, original))
    assert [row["values"] for row in parsed] == [row["values"] for row in original]
    assert all(field["ref"].startswith(sheet + "!") for row in parsed for field in row["fields"].values())

def test_official_xlsx_is_nine_sheet_empty_template():
    from ingest.catalog_template import _xlsx_tables
    content = Path("backend/ingest/templates/catalogo-v1.xlsx").read_bytes()
    assert {name for name, _ in _xlsx_tables(content)} == set(SCHEMAS)
    assert parse_structured("XLSX", content) == []


def test_filled_official_xlsx_keeps_blank_cells_unknown():
    from io import BytesIO, StringIO
    import csv
    from openpyxl import load_workbook
    workbook=load_workbook("backend/ingest/templates/catalogo-v1.xlsx")
    original=sorted(catalog_rows(),key=lambda row:list(SCHEMAS).index(row["sheet"]))
    for sheet in SCHEMAS:
        table=list(csv.reader(StringIO(export_csv(sheet,original).decode("utf-8-sig")),delimiter=";"))
        for index, row in enumerate(table[2:],5):
            for column, value in enumerate(row,1):
                workbook[sheet].cell(index,column,value)
    output=BytesIO()
    workbook.save(output)
    parsed=parse_structured("XLSX",output.getvalue())
    assert len(parsed)==len(original)
    assert [row["values"] for row in parsed]==[row["values"] for row in original]

def test_exact_xlsx_numeric_lexeme_and_formula_rejection():
    from io import BytesIO
    from zipfile import ZipFile, ZIP_DEFLATED
    from xml.etree import ElementTree as ET
    from ingest.catalog_template import _xlsx_tables
    content = Path("backend/ingest/templates/catalogo-v1.xlsx").read_bytes()
    stream = BytesIO()
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(BytesIO(content)) as original, ZipFile(stream, "w", ZIP_DEFLATED) as modified:
        for name in original.namelist():
            value = original.read(name)
            if name == "xl/worksheets/sheet2.xml":
                root = ET.fromstring(value)
                for existing in root.find(ns + "sheetData").findall(ns + "row"):
                    if existing.attrib["r"] == "14":
                        root.find(ns + "sheetData").remove(existing)
                row = ET.SubElement(root.find(ns + "sheetData"), ns + "row", {"r": "14"})
                cell = ET.SubElement(row, ns + "c", {"r": "F14"})
                ET.SubElement(cell, ns + "v").text = "0.1000000000000000001"
                formula = ET.SubElement(row, ns + "c", {"r": "G14"})
                ET.SubElement(formula, ns + "f").text = "3000+3000"
                ET.SubElement(formula, ns + "v").text = "6000"
                value = ET.tostring(root)
            modified.writestr(name, value)
    tables = _xlsx_tables(stream.getvalue())
    numeric = next(row for name, rows in tables if name == "Perfiles" for row in rows if row["_row"] == "14")
    assert numeric["F"] == "0.1000000000000000001"
    assert numeric["G"].startswith("=FORMULA")
    parsed = parse_structured("XLSX", stream.getvalue())
    assert any("Largo comercial" in error["message"] for row in parsed for error in row["errors"])

def test_row_error_names_row_column_and_action():
    row = candidate("Roles y reglas de corte", {"system_code": "SERIE", "sku": "MARCO", "role": "Marco", "welding_loss_per_end_mm": "seis"}, key="bad", row=14, method="MANUAL")
    error = next(error for error in row["errors"] if error["field"] == "welding_loss_per_end_mm")
    assert error["message"] == "Fila 14, columna ‘Pérdida de soldadura’: debe ser un número en mm."

@pytest.mark.parametrize("value", [0.1, float("nan"), True, "-1", "=A1*2", "Infinity"])
def test_manual_numeric_never_accepts_inexact_or_formula(value):
    column = next(column for column in SCHEMAS["Perfiles"] if column.key == "face_width_mm")
    with pytest.raises(ValueError):
        normalize(value, column)

def ai_response(value="60.00", confidence="HIGH", literal="60,00", ref="página 1"):
    return {"records": [{"sheet": "Perfiles", "values": {"system_code": "S60", "sku": "MARCO", "face_width_mm": value}, "evidence": {
        "system_code": {"ref": ref, "quote": "Sistema S60", "literal": "S60", "confidence": "HIGH"},
        "sku": {"ref": ref, "quote": "SKU MARCO", "literal": "MARCO", "confidence": "HIGH"},
        "face_width_mm": {"ref": ref, "quote": "Ancho de cara 60,00 mm", "literal": literal, "confidence": confidence},
    }}]}
@pytest.mark.parametrize("value,confidence,literal,ref,expected", [
    ("60.00", "HIGH", "60,00", "página 1", "60.00"),
    ("600.00", "HIGH", "60,00", "página 1", None),
    ("60.00", "LOW", "60,00", "página 1", None),
    ("60.00", "HIGH", "60,00", "página inventada", None),
    ("0.00", "HIGH", "0", "página 1", None),
])
def test_ai_verified_literal_only_and_low_is_unknown(value, confidence, literal, ref, expected):
    rows = parse_response(json.dumps(ai_response(value, confidence, literal, ref)), tagged=[("Sistema S60. SKU MARCO. Ancho de cara 60,00 mm.", "página 1")], file_name="ficha.pdf")
    assert rows[0]["values"]["face_width_mm"] == expected
    assert rows[0]["values"]["weight_kg_m"] is None
    assert rows[0]["values"]["source"] == "ficha.pdf"

def test_scanned_source_never_promotes_provider_confidence():
    rows = parse_response(ai_response(), tagged=[], file_name="foto.png")
    assert rows[0]["values"]["face_width_mm"] is None
    assert rows[0]["fields"]["face_width_mm"]["proposed"] == "60.00"

def test_preview_new_catalog_all_types_is_stable_and_rejects_duplicates(monkeypatch):
    monkeypatch.setattr(catalog_review, "rows", lambda sql, params: [{"allowed": True}] if "pricing_role" in sql else [])
    original = catalog_rows()
    assert not any(row["errors"] for row in original)
    imported = {"id": uuid4()}
    changes, errors, token = catalog_review._plan(uuid4(), imported, deepcopy(original))
    assert errors == []
    assert len(changes) == len(original)+len([row for row in original if row["sheet"]=="Vidrios"])
    assert {"profile_systems", "profile_articles", "glazing_bead_matrix", "hardware_kits", "glass_purchase_mappings", "catalog_color_skus", "cost_list_items"} == {change.table for change in changes}
    assert token == catalog_review._plan(uuid4(), imported, deepcopy(original))[2]
    duplicate = deepcopy(original[0])
    duplicate["key"] = "duplicate"
    _, errors, _ = catalog_review._plan(uuid4(), imported, deepcopy(original) + [duplicate])
    assert any("duplicada" in error["message"] for error in errors)

def test_review_rejects_foreign_rows_and_fields():
    saved = catalog_rows()[0]
    imported = {"candidates": [saved]}
    for item in ({"key": "foreign", "values": {}}, {"key": saved["key"], "values": {"legacy_authority": True}}):
        with pytest.raises(ContractAPIException):
            catalog_review._entries(imported, [item])

def test_review_correction_preserves_original_source():
    saved = catalog_rows()[0]
    items = [{"key": saved["key"], "values": {**saved["values"], "depth_mm": "62.25"}}]
    reviewed = catalog_review._entries({"candidates": [saved]}, items)[0]
    assert reviewed["values"]["depth_mm"] == "62.25"
    assert reviewed["fields"]["depth_mm"]["method"] == "HUMAN_CORRECTION"
    assert reviewed["fields"]["depth_mm"]["quote"] == saved["fields"]["depth_mm"]["quote"]
    assert reviewed["fields"]["depth_mm"]["original"] == "60.00"

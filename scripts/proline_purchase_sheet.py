"""Fill the official workbook with how a workshop buys Proline PRO6004.

The series itself (profiles, bead, rules, openings and Vorne kits) comes from
the reference library. This sheet adds what the price consumes: the plain bar
purchase of each profile, the steel bars and the double glazing, with the
codes printed in the owner's work order N° 385. Prices are not here; they
belong to the workshop's supplier list.
"""

from __future__ import annotations

import argparse
import csv
from io import StringIO
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "backend"), str(ROOT / "engine/src")]

SOURCE = "OT 385 del taller (NuveraPro, 29-10-2025), págs. 1, 6 y 7"
PROFILES = ("PRO01", "PRO521", "PRO03", "PRO05")
STEELS = (
    ("PRO01", "PRO-R1", "Refuerzo múltiple Pro-R1 28×25×28 · 1,3 mm"),
    ("PRO521", "PRO-R2", "Refuerzo múltiple Pro-R2 29×28×18 · 1,3 mm"),
    ("PRO03", "PRO-R1", "Refuerzo múltiple Pro-R1 28×25×28 · 1,3 mm"),
)


def rows(code: str) -> list[tuple[str, dict]]:
    out: list[tuple[str, dict]] = []
    for sku in PROFILES:
        out.append(("Colores y SKU por color", {
            "system_code": code, "sku": sku, "finish": "WHITE", "commercial_sku": sku,
            "physical_stock_identity": sku, "manufacturer_name": "Proline", "purchase_unit": "BAR",
            "source": SOURCE}))
    for parent, steel, name in STEELS:
        out.append(("Barras de refuerzo", {
            "system_code": code, "sku": parent, "reinforcement_sku": steel, "commercial_sku": steel,
            "name": name, "stock_length_mm": "5800.00", "purchase_unit": "BAR",
            "manufacturer_name": "Proline", "source": SOURCE}))
    out.append(("Fijaciones", {
        "system_code": code, "sku": "120500051", "name": "Tornillo de refuerzo 3,9 × 19 mm",
        "purchasing_sku": "120500051", "manufacturer_name": "Proveedor de herrajes del taller",
        "version": 1, "source": SOURCE}))
    out.append(("Vidrios", {
        "system_code": code, "sku": "DVH-4-12-4", "glass_spec": "4-12-4 Float Incoloro",
        "name": "DVH 4 incoloro + 12 + 4 incoloro", "purchasing_sku": "DVH-4-12-4",
        "manufacturer_name": "Proveedor de vidrio del taller", "version": 1,
        "glass_thickness_mm": "20.00", "bead_sku": "PRO05", "bead_width_mm": "19.00",
        "gasket_interior_mm": "3.00", "gasket_exterior_mm": "3.00", "cut_add_mm": "10.00",
        # Two 4 mm float panes at 2.5 kg/m² per mm; the area limit is the
        # DVH 4 + 4 reference the workshop's supplier sheet would state.
        "weight_kg_m2": "20.00", "max_area_m2": "2.60",
        "source": SOURCE + "; peso y área máxima: ficha de vidrio del taller de ensayo"}))
    return out


def main() -> None:
    from ingest.catalog_template import SCHEMAS, candidate, export_csv, parse_structured

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--system-code", default="PRO6004")
    args = parser.parse_args()
    entries = [candidate(sheet, values, key=f"r{index}", row=index + 5, method="MANUAL")
               for index, (sheet, values) in enumerate(rows(args.system_code))]
    assert not any(entry["errors"] for entry in entries), [e["errors"] for e in entries if e["errors"]]
    book = load_workbook(ROOT / "backend/ingest/templates/catalogo-v1.xlsx")
    for name in SCHEMAS:
        table = list(csv.reader(StringIO(export_csv(name, entries).decode("utf-8-sig")), delimiter=";"))
        for index, values in enumerate(table[2:], 5):
            for column, value in enumerate(values, 1):
                book[name].cell(index, column, value)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)
    parsed = parse_structured("XLSX", args.output.read_bytes())
    assert parsed and len(parsed) == len(entries) and not any(row["errors"] for row in parsed)
    print(f"{len(parsed)} filas de compra Proline preparadas.")


if __name__ == "__main__":
    main()

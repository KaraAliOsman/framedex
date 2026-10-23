"""Fill the official workbook with exact synthetic data for repeatable browser QA."""
from __future__ import annotations

import argparse
import csv
from io import StringIO
from pathlib import Path
import sys
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT),str(ROOT/'backend'),str(ROOT/'engine/src')]


def main():
    from backend.tests.catalog_interchange_fixture import catalog_rows
    from ingest.catalog_template import SCHEMAS, export_csv, parse_structured
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--system-code',default='REVISION-D01')
    parser.add_argument('--openings',action='store_true',help='Include D03 sourced physical-opening capabilities')
    parser.add_argument('--system-name',help='Human name for an isolated synthetic browser review')
    parser.add_argument('--bar-purchases',action='store_true',
        help='Declare a white bar purchase per profile under its own SKU (a quotable series)')
    parser.add_argument('--rename',action='append',default=[],metavar='OLD=NEW',
        help='Relabel synthetic text in every cell, e.g. DEMO_60=PRO6004-ENSAYO')
    args=parser.parse_args()
    if not args.system_code or len(args.system_code)>50:
        parser.error('system code must contain 1–50 characters')
    book=load_workbook(ROOT/'backend/ingest/templates/catalogo-v1.xlsx')
    entries=catalog_rows()
    for entry in entries:
        entry['values']['system_code']=args.system_code
        if args.openings and entry['sheet']=='Vidrios':
            entry['values']['sku']=f"{args.system_code}-{entry['values']['sku']}"
            entry['values']['purchasing_sku']=entry['values']['sku']
        if entry['sheet']=='Sistemas':
            entry['values']['name']='Catálogo de revisión sintético D01'
            if args.openings:
                from backend.catalogs.demo_openings import opening_manifest
                capabilities=next(record['params']['opening_capabilities'] for record in opening_manifest() if record['code']=='DEMO_60')
                entry['values'].update(name='DEMO · Aperturas revisadas D03',opening_capabilities=[cap for cap in capabilities
                    if cap['leaf_role']=='SINGLE' and cap['movement'] in ('FIXED','TURN') and not cap['fixed_in_sash']],paired_leaf_rule=None)
            if args.system_name:
                entry['values']['name']=args.system_name
    if args.bar_purchases:
        from ingest.catalog_template import candidate
        colors=[entry for entry in entries if entry['sheet']=='Colores y SKU por color']
        template=dict(colors[0]['values'])
        entries=[entry for entry in entries if entry['sheet']!='Colores y SKU por color']
        profiles=[entry['values']['sku'] for entry in entries if entry['sheet']=='Perfiles']
        last=max(entry['row'] for entry in entries)
        for index,sku in enumerate(profiles,1):
            entries.append(candidate('Colores y SKU por color',{**template,'sku':sku,'commercial_sku':sku,
                'physical_stock_identity':sku,'stock_color':'WHITE','purchase_unit':'BAR',
                'manufacturer_name':'Fabricante de ensayo','cutting_profile_code':'DEMO','binding_version':1},
                key=f'bar{index}',row=last+index,method='MANUAL'))
        # Color purchases count only under a finish chart: a white-only synthetic one.
        from backend.catalogs.demo_finishes import finish_manifest
        chart=next(record['params']['finish_authority'] for record in finish_manifest() if record['code']=='DEMO_60')
        chart={**chart,'colors':[color for color in chart['colors'] if color['code']=='WHITE'],
            'combinations':[combo for combo in chart['combinations'] if combo['code']=='WHITE']}
        for entry in entries:
            if entry['sheet']=='Sistemas':
                entry['values'].update(finish_authority=chart,finishes=['WHITE'])
            if entry['sheet']=='Barras de refuerzo':
                entry['values'].update(stock_color='WHITE',cutting_profile_code='DEMO',binding_version=1)
    renames=[item.split('=',1) for item in args.rename]
    if any(len(pair)!=2 or not pair[0] for pair in renames):
        parser.error('--rename takes OLD=NEW')
    def relabel(value):
        if isinstance(value,str):
            for old,new in renames:
                value=value.replace(old,new)
        elif isinstance(value,dict):
            return {key:relabel(item) for key,item in value.items()}
        elif isinstance(value,list):
            return [relabel(item) for item in value]
        return value
    for entry in entries:
        entry['values']=relabel(entry['values'])
    for name in SCHEMAS:
        table=list(csv.reader(StringIO(export_csv(name,entries).decode('utf-8-sig')),delimiter=';'))
        for index,values in enumerate(table[2:],5):
            for column,value in enumerate(values,1):
                book[name].cell(index,column,value)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    book.save(args.output)
    parsed=parse_structured('XLSX',args.output.read_bytes())
    assert parsed and len(parsed)==len(entries) and not any(row['errors'] for row in parsed)
    print(f'{len(parsed)} exact synthetic rows prepared.')


if __name__=='__main__':
    main()

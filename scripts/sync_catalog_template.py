"""Refresh the official blank workbook against the authoritative import schema."""
from copy import copy
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.cell_range import CellRange

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT / "engine/src"))
from ingest.catalog_template import SCHEMAS  # noqa: E402


def main() -> None:
    path = ROOT / "backend/ingest/templates/catalogo-v1.xlsx"
    book = load_workbook(path)
    for name, columns in SCHEMAS.items():
        sheet = book[name]
        if any(cell.value is not None for row in sheet.iter_rows(min_row=5) for cell in row):
            raise ValueError("Refusing to overwrite a populated catalog template")
        styles = [copy(sheet.cell(row, 1)._style) for row in range(1, 25)]
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        sheet.data_validations.dataValidation.clear()
        for index, column in enumerate(columns, 1):
            letter = get_column_letter(index)
            sheet.cell(4, index, column.label)
            for row in range(4, 25):
                sheet.cell(row, index)._style = copy(styles[row - 1])
                sheet.cell(row, index).number_format = "@"
            sheet.column_dimensions[letter].width = min(48, max(22, len(column.label) + 2))
            labels = list(column.choices.values()) if column.choices else ["Sí", "No"] if column.kind == "boolean" else []
            if labels:
                validation = DataValidation(type="list", formula1='"' + ",".join(labels) + '"', allow_blank=not column.required)
                validation.errorTitle = "Revise la opción"
                validation.error = "Elija un valor de la lista de esta columna."
                validation.showErrorMessage = True
                sheet.add_data_validation(validation)
                validation.add(CellRange(f"{letter}5:{letter}24"))
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
        sheet.auto_filter.ref = f"A4:{get_column_letter(len(columns))}24"
    book.save(path)


if __name__ == "__main__":
    main()
